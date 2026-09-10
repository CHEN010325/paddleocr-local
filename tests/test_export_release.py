"""Export degraded OCR, formula fallbacks and mixed-page release scenarios."""
import base64
import io

import pytest
from docx import Document
from lxml import etree
from PIL import Image
from pypdf import PdfReader

import exporters as ex


def test_mixed_document_preserves_footnotes_and_landscape_transitions():
    headers = ''.join(f'<th>Column {n}</th>' for n in range(10))
    cells = ''.join(f'<td>{n}</td>' for n in range(10))
    table = f'<table><tr>{headers}</tr><tr>{cells}</tr></table>'
    markdown = f'# Report\n\n{table}\n\nFootnote: source attribution.\n\n{table}\n\n```\ncode\n```\n'
    document = Document(io.BytesIO(ex.build_docx({'markdown': markdown})))
    assert len(document.tables) == 2
    assert len(document.sections) == 5
    assert 'Footnote: source attribution.' in '\n'.join(p.text for p in document.paragraphs)
    assert document.sections[-1].page_width < document.sections[-1].page_height


@pytest.mark.parametrize('formula', [
    r'\begin{cases}a&x>0\\b&x<0\end{cases}',
    r'\begin{array}{ll}\hline\\a&\\\hline\end{array}+b',
    r'\unknowncommand{x}',
    r'\begin{matrix}a&b\\c&d\end{matrix}',
    r'\begin{equation}x^2\tag{9}\end{equation}',
])
def test_ocr_formula_variants_remain_visible(formula):
    doc = Document(io.BytesIO(ex.build_docx({'markdown': formula})))
    xml = doc._element.xml
    assert doc.inline_shapes or 'oMath' in xml or formula in '\n'.join(p.text for p in doc.paragraphs)


def test_formula_fallback_when_rasterizer_is_missing(monkeypatch):
    monkeypatch.setattr(ex, '_load_formula_rasterizer', lambda: (None, None))
    image = ex._render_mathtext_piece(r'\alpha+\unknowncommand{x}')
    assert image.width > 8 and image.getchannel('A').getbbox()
    monkeypatch.setattr(ex, '_formula_image', lambda *args, **kwargs: None)
    monkeypatch.setattr(ex, '_formula_omml', lambda *args, **kwargs: None)
    doc = Document()
    paragraph = doc.add_paragraph()
    assert ex._append_formula(paragraph, r'x+y\tag{3}')
    assert 'x' in paragraph.text and 'y' in paragraph.text
    assert not ex._append_formula(doc.add_paragraph(), '')


def test_formula_fallback_keeps_editable_math_when_image_fails(monkeypatch):
    monkeypatch.setattr(ex, '_formula_image', lambda *args, **kwargs: b'broken image')
    doc = Document()
    assert ex._append_formula(doc.add_paragraph(), r'x^2\tag{4}')
    assert 'oMath' in doc._element.xml
    assert '(4)' in doc.paragraphs[0].text


def test_native_math_falls_back_if_converter_fails(monkeypatch):
    monkeypatch.setenv('PANDOCR_DOCX_FORMULA_MODE', 'native')
    monkeypatch.setattr(ex, '_latex_to_mathml', lambda _: (_ for _ in ()).throw(ValueError('invalid formula')))
    assert ex._formula_omml('x') is None
    doc = Document(io.BytesIO(ex.build_docx({'markdown': '$$x^2\\tag{7}$$'})))
    assert doc.inline_shapes and '(7)' in doc._element.xml


def test_native_math_rejects_invalid_converter_output(monkeypatch):
    monkeypatch.setattr(ex, '_latex_to_mathml', lambda value: value)
    monkeypatch.setattr(ex, '_mathml_to_omml', lambda _: '')
    assert ex._formula_omml('x') is None
    monkeypatch.setattr(ex, '_mathml_to_omml', lambda _: '<m:oMath><m:r><m:t>x</m:t></m:r></m:oMath>')
    assert ex._formula_omml('x') is not None


def test_html_ocr_tables_tolerate_bad_spans_and_preserve_cell_text():
    html = '<table><tr><td rowspan="oops" colspan="bad">First<br>Second</td><td>42</td></tr><tr></tr></table>'
    specs = ex.extract_html_table_specs(html)
    assert specs[0].rows == [['First\nSecond', '42']]
    assert ex.extract_html_table_specs('<table><tr></tr></table>') == []
    doc = Document(io.BytesIO(ex.build_docx({'markdown': html + '\n\n<img src="missing.png" alt="Footnote figure" width="300px">'})))
    assert 'First' in doc.tables[0].cell(0, 0).text
    assert 'Footnote figure' in '\n'.join(p.text for p in doc.paragraphs)


def test_image_aliases_resolve_across_batches_without_network():
    out = io.BytesIO()
    Image.new('RGB', (12, 10), 'blue').save(out, 'PNG')
    encoded = base64.b64encode(out.getvalue()).decode()
    task = {'markdown': '<img src="imgs/figure.png" width="40%" alt="Figure">',
            'images': {'ocr_images/batch_figure.png': encoded},
            'ocrResults': [None, {'markdown': {'images': {'./imgs/figure.png': 'ocr_images/batch_figure.png'}}}],
            'batchMarkdown': {'1': {'markdown': {'images': {'imgs/figure.png': 'ocr_images/batch_figure.png'}}}}}
    assert ex._image_bytes(task, 'imgs/figure.png') == out.getvalue()
    doc = Document(io.BytesIO(ex.build_docx(task)))
    assert len(doc.inline_shapes) == 1
    assert ex._parse_html_image_line('<img alt="missing">') is None
    assert ex._parse_html_image_line('<img src="x" width="1191px">')[2] == 6.25


def test_searchable_pdf_retains_raw_footnote_blocks(tmp_path):
    from reportlab.pdfgen import canvas
    source = tmp_path / 'source.pdf'
    pdf = canvas.Canvas(str(source), pagesize=(200, 300))
    pdf.showPage()
    pdf.save()
    task = {'mimeType': 'application/pdf', 'ocrResults': [
        {'sourcePage': 1, 'parsing_res_list': [None, {'block_label': 'footnote', 'block_content': 'Retained footnote 12345'}]}
    ]}
    result = ex.build_searchable_pdf(task, source)
    assert 'Retained footnote 12345' in PdfReader(io.BytesIO(result)).pages[0].extract_text()


def test_formula_text_rendering_survives_missing_system_fonts(monkeypatch):
    monkeypatch.setattr(ex.Path, 'exists', lambda _: False)
    image = ex._render_formula_text_image('x + y')
    assert image.width > 8 and image.height > 8


def test_math_renderer_load_failure_does_not_break_exports(monkeypatch):
    import builtins
    original = builtins.__import__
    def unavailable(name, *args, **kwargs):
        if name == 'matplotlib':
            raise ImportError('optional renderer absent')
        return original(name, *args, **kwargs)
    monkeypatch.setattr(ex, '_MATH_RENDERER_READY', False)
    monkeypatch.setattr(ex, '_MATH_TO_IMAGE', None)
    monkeypatch.setattr(ex, '_MATH_FONT_PROPERTIES', None)
    monkeypatch.setattr(builtins, '__import__', unavailable)
    assert ex._load_formula_rasterizer() == (None, None)


def test_empty_formula_and_escaped_alignment_cells():
    assert ex._formula_image('') is None
    assert ex._compose_array_formula(r'\begin{array}{l}\hline\end{array}') is None
    assert ex._split_formula_cells(r'a\&b&c') == [r'a\&b', 'c']
    assert ex._shorthand_formula_source('ordinary prose') == ''
    assert ex._shorthand_formula_source(r'ordinary \arrayll a&b') == ''
    assert ex._shorthand_formula_source(r'$x \arrayll a&b') == ''
    assert ex._looks_like_formula('') is False
    # Escaped equation delimiters and accidental currency must stay literal.
    assert ex.plain_markdown_text(r'Price \$5 and \$10') == 'Price $5 and $10'


def test_searchable_pdf_rejects_password_protected_source(tmp_path):
    from pypdf import PdfWriter
    source = tmp_path / 'encrypted.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=300)
    writer.encrypt('secret')
    with source.open('wb') as stream:
        writer.write(stream)
    with pytest.raises(ex.ExportError, match='encrypted'):
        ex.build_searchable_pdf({'markdown':'text', 'mimeType':'application/pdf'}, source)


def test_partial_page_text_preserves_page_count(tmp_path):
    from pypdf import PdfWriter
    source = tmp_path / 'two.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=300)
    writer.add_blank_page(width=200, height=300)
    with source.open('wb') as stream:
        writer.write(stream)
    task = {'ocrResults':[{'sourcePage':2, 'text':'Second page only'}], 'mimeType':'application/pdf'}
    reader = PdfReader(io.BytesIO(ex.build_searchable_pdf(task, source)))
    assert len(reader.pages) == 2 and 'Second page only' in reader.pages[1].extract_text()
    assert not reader.pages[0].extract_text()


def test_missing_or_corrupt_export_images_preserve_captions(monkeypatch):
    document = Document(io.BytesIO(ex.build_docx({'markdown': '![Caption](missing.png)\n\n<img src="bad.png" alt="Other">',
                                                 'images': {'bad.png': 'bad base64'}})))
    assert 'Caption' in document._element.xml and 'Other' in document._element.xml
    monkeypatch.setattr(ex, '_image_bytes', lambda *args: b'bad image')
    doc = Document()
    assert not ex._add_markdown_image(doc, {}, 'caption', 'bad')
    assert not doc.paragraphs


def test_ocr_algorithm_and_adjacent_wide_tables_have_no_spurious_sections():
    wide = '<table><tr>' + '<td>Value</td>' * 10 + '</tr></table>'
    doc = Document(io.BytesIO(ex.build_docx({'markdown':
        'Algorithm 1 OCR\n\nInput: image\nRead text\n\nFootnote after algorithm\n\n' + wide + '\n\n' + wide})))
    assert len(doc.sections) == 2
    assert len(doc.tables) == 2
    assert 'Footnote after algorithm' in doc._element.xml


def test_omml_converter_namespace_wrappers_are_supported(monkeypatch):
    monkeypatch.setattr(ex, '_latex_to_mathml', lambda value: value)
    monkeypatch.setattr(ex, '_mathml_to_omml', lambda _: '<m:oMath attr="x"><m:r><m:t>x</m:t></m:r></m:oMath>')
    assert ex._formula_omml('x') is not None
    monkeypatch.setattr(ex, '_mathml_to_omml', lambda _: '<m:root xmlns:m="' + ex.MATH_NS + '"><m:oMath/></m:root>')
    assert ex._formula_omml('x') is not None
    monkeypatch.setattr(ex, '_latex_to_mathml', None)
    assert ex._formula_omml('x') is None


def test_native_equation_tag_is_preserved(monkeypatch):
    monkeypatch.setenv('PANDOCR_DOCX_FORMULA_MODE', 'native')
    doc = Document()
    assert ex._append_formula(doc.add_paragraph(), r'x+y\tag{12}', display=True)
    assert '(12)' in doc.paragraphs[0].text
    assert 'oMathPara' in doc._element.xml


def test_large_text_overlay_and_empty_html_table():
    reader = PdfReader(io.BytesIO(ex._invisible_overlay((72, 12), ['x' * 2000, 'tail'])))
    assert 'tail' in reader.pages[0].extract_text()
    assert ex._layout_html_table([]) is None
    assert ex._layout_html_table([[]]) is None
    doc = Document(io.BytesIO(ex.build_docx({'markdown':'<table></table>\n\nStill readable'})))
    assert 'Still readable' in doc._element.xml


def test_searchable_pdf_reports_decryption_errors(tmp_path, monkeypatch):
    source = tmp_path / 'encrypted.pdf'
    source.write_bytes(b'%PDF-')
    class BrokenReader:
        is_encrypted = True
        def decrypt(self, password):
            raise ValueError('bad encryption metadata')
    monkeypatch.setattr(ex, 'PdfReader', lambda _: BrokenReader())
    with pytest.raises(ex.ExportError, match='encrypted'):
        ex.build_searchable_pdf({'markdown':'text', 'mimeType':'application/pdf'}, source)


def test_formula_fallback_skips_unreadable_font(monkeypatch):
    from PIL import ImageFont
    original = ImageFont.truetype
    calls = []
    def broken_first(*args, **kwargs):
        calls.append(True)
        if len(calls) == 1:
            raise OSError('font unavailable')
        return original(*args, **kwargs)
    monkeypatch.setattr(ImageFont, 'truetype', broken_first)
    assert ex._render_formula_text_image('x+y').width > 8
